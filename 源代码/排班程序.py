import csv, sys, os, logging, traceback, webbrowser, shutil
from pathlib import Path
from datetime import datetime, date, time, timedelta
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from tkinter.scrolledtext import ScrolledText
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from ortools.sat.python import cp_model

PROJECT = Path(__file__).resolve().parent.parent
ASSETS = Path(sys._MEIPASS) / '内置资料' if getattr(sys, 'frozen', False) else PROJECT
ROOT = Path.home() / '排班小程序数据'
REPO_URL = 'https://github.com/fhhv5/dorm-duty-scheduler'
(ROOT/'日志').mkdir(parents=True, exist_ok=True)
(ROOT/'输出').mkdir(parents=True, exist_ok=True)
logging.basicConfig(filename=ROOT/'日志'/'运行日志.log', level=logging.INFO, format='%(asctime)s %(levelname)s %(message)s', encoding='utf-8')

def read_table(path, required):
    if not path: raise ValueError('请选择所有输入文件')
    p=Path(path)
    if p.suffix.lower()=='.csv':
        with open(p, encoding='utf-8-sig', newline='') as f: rows=list(csv.reader(f))
    elif p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, read_only=True, data_only=True)
        try: rows=list((wb['数据'] if '数据' in wb.sheetnames else wb.active).values)
        finally: wb.close()
    else: raise ValueError('仅支持 .xlsx 或 UTF-8 CSV 文件')
    if not rows: raise ValueError(f'{p.name} 是空文件')
    headers=[str(v).strip() if v is not None else '' for v in rows[0]]
    if len(headers)!=len(set(headers)): raise ValueError(f'{p.name} 表头重复')
    missing=set(required)-set(headers)
    if missing: raise ValueError(f'{p.name} 缺少表头：{",".join(missing)}')
    return [(i,dict(zip(headers,row))) for i,row in enumerate(rows[1:],2) if any(v is not None and str(v).strip() for v in row)]

def s(v): return str(v).strip() if v is not None else ''
def dt(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    return datetime.strptime(s(v).replace('/','-').replace('.','-'),'%Y-%m-%d').date()
def tm(v):
    if isinstance(v,datetime): return v.time().replace(second=0,microsecond=0)
    if isinstance(v,time): return v.replace(second=0,microsecond=0)
    if isinstance(v,(int,float)): return (datetime.min+timedelta(days=float(v))).time().replace(second=0,microsecond=0)
    return datetime.strptime(s(v).replace('：',':'),'%H:%M').time()

TEMPLATES = {'名单': '排班名单模板.xlsx', '班次': '班次设置模板.xlsx', '冲突': '时间冲突模板.xlsx'}

def export_template(kind, destination):
    if kind not in TEMPLATES: raise ValueError('未知的模板类型')
    shutil.copyfile(ASSETS / '模板' / TEMPLATES[kind], destination)

def schedule(peoplefile, shiftsfile, conflictsfile, out, include_phone=True, include_dorm=True, include_signature=True, include_class=True, include_major=True, include_counselor=True):
    people=[]
    for i,r in read_table(peoplefile,['姓名']):
        name=s(r['姓名'])
        if not name: raise ValueError(f'名单第 {i} 行姓名为空')
        if any(p['姓名']==name for p in people): raise ValueError(f'姓名重复：{name}，请使用唯一姓名')
        people.append({'姓名':name,'电话':s(r.get('电话')),'宿舍':s(r.get('宿舍')),'班级':s(r.get('班级')),'专业':s(r.get('专业')),'辅导员':s(r.get('辅导员'))})
    if not people: raise ValueError('名单没有成员')
    shifts=[]
    for i,r in read_table(shiftsfile,['班次编号','日期','班次名称','开始时间','结束时间','人数']):
        try:
            ident=s(r['班次编号']); d=dt(r['日期']); start=datetime.combine(d,tm(r['开始时间'])); end=datetime.combine(d,tm(r['结束时间']))
            if end<=start: end+=timedelta(days=1)
            n=int(s(r['人数']))
            if not ident or not s(r['班次名称']) or n<1 or n>len(people): raise ValueError('编号、名称或人数不合法')
            if any(x['编号']==ident for x in shifts): raise ValueError('班次编号重复')
            shifts.append({'编号':ident,'名称':s(r['班次名称']),'开始':start,'结束':end,'人数':n})
        except Exception as e: raise ValueError(f'班次表第 {i} 行：{e}')
    if not shifts: raise ValueError('班次表没有班次')
    blocked=set(); names={p['姓名'] for p in people}; ids={x['编号'] for x in shifts}
    for i,r in read_table(conflictsfile,['姓名','冲突班次编号']):
        name=s(r['姓名']); raw=s(r['冲突班次编号']).replace('，',',').replace('、',',').replace('；',',').replace(';',',')
        if name not in names: raise ValueError(f'冲突表第 {i} 行姓名不在名单中：{name}')
        codes=[v.strip() for v in raw.split(',') if v.strip()]
        if not codes: raise ValueError(f'冲突表第 {i} 行班次编号为空')
        for code in codes:
            if code not in ids: raise ValueError(f'冲突表第 {i} 行不存在班次编号：{code}')
            blocked.add((name,code))
    m=cp_model.CpModel(); x={}
    for p in range(len(people)):
        for j,shift in enumerate(shifts):
            x[p,j]=m.NewBoolVar(f'x_{p}_{j}')
            if (people[p]['姓名'],shift['编号']) in blocked: m.Add(x[p,j]==0)
    for j,shift in enumerate(shifts): m.Add(sum(x[p,j] for p in range(len(people)))==shift['人数'])
    for j,a in enumerate(shifts):
        for k in range(j+1,len(shifts)):
            b=shifts[k]
            if a['开始']<b['结束'] and b['开始']<a['结束']:
                for p in range(len(people)): m.Add(x[p,j]+x[p,k]<=1)
    counts=[sum(x[p,j] for j in range(len(shifts))) for p in range(len(people))]
    hi=m.NewIntVar(0,len(shifts),'max'); lo=m.NewIntVar(0,len(shifts),'min')
    m.AddMaxEquality(hi,counts); m.AddMinEquality(lo,counts); m.Minimize(hi-lo)
    solver=cp_model.CpSolver(); solver.parameters.max_time_in_seconds=30
    status=solver.Solve(m)
    if status not in (cp_model.OPTIMAL,cp_model.FEASIBLE):
        detail='；'.join(f"{q['编号']} 可用 {sum((p['姓名'],q['编号']) not in blocked for p in people)} / 需要 {q['人数']}" for q in shifts if sum((p['姓名'],q['编号']) not in blocked for p in people)<q['人数'])
        raise ValueError(('求解超时，未找到结果' if status==cp_model.UNKNOWN else '无法满足全部班次和冲突限制') + ('：'+detail if detail else '；请检查重叠班次和人员冲突，或增加人员。'))
    wb=Workbook(); ws=wb.active; ws.title='排班表'; ws.append(['班次编号','日期','班次名称','开始时间','结束时间','姓名','电话','宿舍','班级','专业','辅导员','签名','备注'])
    for j,q in sorted(enumerate(shifts),key=lambda z:z[1]['开始']):
        for p,person in enumerate(people):
            if solver.Value(x[p,j]): ws.append([q['编号'],q['开始'].strftime('%Y-%m-%d'),q['名称'],q['开始'].strftime('%H:%M'),q['结束'].strftime('%Y-%m-%d %H:%M') if q['结束'].date()!=q['开始'].date() else q['结束'].strftime('%H:%M'),person['姓名'],person['电话'],person['宿舍'],person['班级'],person['专业'],person['辅导员'],'',''])
    summary=wb.create_sheet('人员统计'); summary.append(['姓名','排班次数','电话','宿舍','班级','专业','辅导员'])
    for p,person in enumerate(people): summary.append([person['姓名'],solver.Value(counts[p]),person['电话'],person['宿舍'],person['班级'],person['专业'],person['辅导员']])
    for sheet in wb:
        for col in range(sheet.max_column,0,-1):
            header=sheet.cell(1,col).value
            included = {'电话':include_phone, '宿舍':include_dorm, '签名':include_signature, '班级':include_class, '专业':include_major, '辅导员':include_counselor}
            if header in included and not included[header]: sheet.delete_cols(col)
    for sheet in wb:
        sheet.freeze_panes='A2'; sheet.auto_filter.ref=sheet.dimensions
        for c in sheet[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='245B78'); c.alignment=Alignment(horizontal='center')
        for col in sheet.columns:
            letter=col[0].column_letter; sheet.column_dimensions[letter].width=min(28,max(12,max(len(str(c.value or '')) for c in col)*1.4+2))
    wb.save(out); logging.info('生成成功 人员=%s 班次=%s 状态=%s 输出=%s',len(people),len(shifts),solver.StatusName(status),out)
    return solver.StatusName(status)

def main():
    root=tk.Tk(); root.title('排班小程序'); root.geometry('800x590')
    frame=ttk.Frame(root,padding=22); frame.pack(fill='both',expand=True)
    ttk.Label(frame,text='排班小程序',font=('Microsoft YaHei',18,'bold')).grid(row=0,column=0,columnspan=3,sticky='w',pady=(0,12))
    ttk.Label(frame,text='按模板填写名单、班次和冲突；班次编号必须一致。',font=('Microsoft YaHei',10)).grid(row=1,column=0,columnspan=3,sticky='w',pady=(0,15))
    vals=[tk.StringVar() for _ in range(3)]
    for i,label in enumerate(['排班名单','班次设置','时间冲突']):
        ttk.Label(frame,text=label).grid(row=i+2,column=0,sticky='w',pady=9)
        ttk.Entry(frame,textvariable=vals[i],width=68).grid(row=i+2,column=1,padx=10)
        def choose(v=vals[i]):
            path=filedialog.askopenfilename(filetypes=[('表格','*.xlsx *.csv')]);
            if path: v.set(path)
        ttk.Button(frame,text='选择文件',command=choose).grid(row=i+2,column=2)
    opts=[tk.BooleanVar(value=True) for _ in range(6)]
    options=ttk.Frame(frame); options.grid(row=5,column=0,columnspan=3,sticky='w',pady=8)
    ttk.Label(options,text='导出字段：').pack(side='left')
    for label,v in zip(['电话','宿舍','签名栏','班级','专业','辅导员'],opts): ttk.Checkbutton(options,text=label,variable=v).pack(side='left',padx=6)
    def run():
        out=filedialog.asksaveasfilename(initialdir=ROOT/'输出',initialfile='排班表_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.xlsx',defaultextension='.xlsx',filetypes=[('Excel 工作簿','*.xlsx')])
        if not out:return
        try:
            status=schedule(*(v.get() for v in vals),out,*(v.get() for v in opts))
            messagebox.showinfo('排班完成',f'已保存：\n{out}\n'+('已找到公平差距最小的解。' if status=='OPTIMAL' else '已找到可行解，公平性可能仍可改善。'))
        except Exception as e:
            logging.exception('排班失败'); messagebox.showerror('无法生成排班',str(e))
    ttk.Button(frame,text='生成并导出排班表',command=run).grid(row=6,column=1,pady=16,sticky='w')
    ttk.Label(frame,text='支持 Excel（.xlsx）和 UTF-8 CSV；冲突可留空表。').grid(row=7,column=0,columnspan=3,sticky='w')
    def show_text(title, content, parent=root):
        win=tk.Toplevel(parent); win.title(title); win.geometry('700x550')
        panel=ttk.Frame(win,padding=16); panel.pack(fill='both',expand=True)
        text=ScrolledText(panel,wrap='word',font=('Microsoft YaHei',10))
        text.pack(fill='both',expand=True); text.insert('1.0',content); text.configure(state='disabled')
    def about():
        win=tk.Toplevel(root); win.title('关于与感谢'); win.geometry('600x370')
        panel=ttk.Frame(win,padding=20); panel.pack(fill='both',expand=True)
        ttk.Label(panel,text='排班小程序',font=('Microsoft YaHei',16,'bold')).pack(anchor='w')
        ttk.Label(panel,text='项目发布与维护：wufong（GitHub @fhhv5）\n开发协助：OpenAI Codex\n\n感谢 Google OR-Tools、openpyxl、Python 和 PyInstaller 的开发者，\n以及 GitHub 排班项目作者提供的参考思路。',justify='left').pack(anchor='w',pady=15)
        ttk.Button(panel,text='打开 GitHub 项目仓库',command=lambda:webbrowser.open(REPO_URL)).pack(anchor='w')
        ttk.Label(panel,text=REPO_URL).pack(anchor='w',pady=8)
        def licenses():
            choice=tk.Toplevel(win); choice.title('开源许可'); choice.geometry('720x500')
            box=tk.Listbox(choice); box.pack(fill='both',expand=True,padx=12,pady=12)
            files=sorted((ASSETS/'第三方许可').rglob('*'))
            files=[p for p in files if p.is_file()]
            files.insert(0, ASSETS/'LICENSE')
            for path in files: box.insert('end',str(path.relative_to(ASSETS)))
            def open_license():
                selection=box.curselection()
                if selection:
                    path=files[selection[0]]
                    show_text(path.name,path.read_text(encoding='utf-8',errors='replace'),choice)
            ttk.Button(choice,text='查看所选许可',command=open_license).pack(pady=(0,12))
        ttk.Button(panel,text='查看开源许可',command=licenses).pack(anchor='w')
    def export(kind):
        dest = filedialog.asksaveasfilename(title='保存'+TEMPLATES[kind],initialfile=TEMPLATES[kind],defaultextension='.xlsx',filetypes=[('Excel 工作簿','*.xlsx')])
        if not dest: return
        try:
            export_template(kind,dest)
            messagebox.showinfo('导出完成',f'模板已保存到：\n{dest}')
        except Exception as e:
            logging.exception('导出模板失败'); messagebox.showerror('导出失败',str(e))
    buttons=ttk.Frame(frame); buttons.grid(row=8,column=0,columnspan=3,sticky='w',pady=10)
    for kind in TEMPLATES:
        ttk.Button(buttons,text='导出'+kind+'模板',command=lambda k=kind:export(k)).pack(side='left',padx=(0,8))
    more=ttk.Frame(frame); more.grid(row=9,column=0,columnspan=3,sticky='w',pady=4)
    ttk.Button(more,text='使用说明',command=lambda:show_text('使用说明',(ASSETS/'使用说明.txt').read_text(encoding='utf-8'))).pack(side='left',padx=(0,8))
    ttk.Button(more,text='用户协议',command=lambda:show_text('用户协议',(ASSETS/'用户协议.txt').read_text(encoding='utf-8'))).pack(side='left',padx=(0,8))
    ttk.Button(more,text='关于与感谢 / GitHub',command=about).pack(side='left')
    def welcome():
        if (ROOT/'已阅读首次说明').exists(): return
        win=tk.Toplevel(root); win.title('欢迎使用排班小程序'); win.geometry('700x580'); win.transient(root); win.grab_set()
        panel=ttk.Frame(win,padding=16); panel.pack(fill='both',expand=True)
        ttk.Label(panel,text='首次使用教程与用户协议',font=('Microsoft YaHei',15,'bold')).pack(anchor='w',pady=(0,10))
        text=ScrolledText(panel,wrap='word',font=('Microsoft YaHei',10),height=22)
        text.pack(fill='both',expand=True)
        text.insert('1.0',(ASSETS/'使用说明.txt').read_text(encoding='utf-8')+'\n\n'+(ASSETS/'用户协议.txt').read_text(encoding='utf-8'))
        text.configure(state='disabled')
        def start():
            try: (ROOT/'已阅读首次说明').write_text('已查看首次使用说明\n',encoding='utf-8')
            except OSError as e:
                messagebox.showerror('保存失败',str(e),parent=win); return
            win.destroy()
        ttk.Button(panel,text='我已阅读，开始使用',command=start).pack(pady=(12,0))
        win.protocol('WM_DELETE_WINDOW',root.destroy)
    root.after(100,welcome)
    root.mainloop()
if __name__=='__main__':
    if len(sys.argv)==4 and sys.argv[1]=='--export-template':
        export_template(sys.argv[2],sys.argv[3])
    else:
        main()
